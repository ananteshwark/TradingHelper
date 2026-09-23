"""Parsers for BSE, broker (Tier 2) and Screener.in (Tier 3) masters/exports.

STATUS: BSE and Angel One layouts follow their published payloads and have not
yet been checked against a verified sample from this environment.
"""

from __future__ import annotations

import csv
import io
import json

import polars as pl

from igs.normalize.nse import SchemaMismatch, _num, _require


def parse_bse_scrips(content: bytes) -> pl.DataFrame:
    data = json.loads(content)
    if isinstance(data, dict) and "Table" in data:
        data = data["Table"]
    if not isinstance(data, list):
        raise SchemaMismatch("BSE scrip payload is not a list")
    out = []
    for r in data:
        _require(r.keys(), ["SCRIP_CD", "Scrip_Name", "ISIN_NUMBER"], "BSE scrip row")
        out.append({"scrip_code": str(r["SCRIP_CD"]).strip(),
                    "isin": (r.get("ISIN_NUMBER") or "").strip() or None,
                    "scrip_id": (r.get("scrip_id") or "").strip() or None,
                    "name": r["Scrip_Name"].strip(), "status": r.get("Status"),
                    "scrip_group": (r.get("GROUP") or "").strip() or None,
                    "face_value": _num(str(r.get("FACE_VALUE") or "")),
                    "industry": r.get("INDUSTRY")})
    return pl.DataFrame(out, schema={"scrip_code": pl.Utf8, "isin": pl.Utf8,
                                     "scrip_id": pl.Utf8, "name": pl.Utf8, "status": pl.Utf8,
                                     "scrip_group": pl.Utf8, "face_value": pl.Float64,
                                     "industry": pl.Utf8})


def parse_angel_master(content: bytes) -> pl.DataFrame:
    """Angel One instrument master; keeps NSE/BSE cash-equity lines only."""
    data = json.loads(content)
    if not isinstance(data, list):
        raise SchemaMismatch("Angel master is not a list")
    out = []
    for r in data:
        _require(r.keys(), ["token", "symbol", "name", "exch_seg", "instrumenttype"],
                 "Angel instrument")
        if r["exch_seg"] not in ("NSE", "BSE") or r["instrumenttype"] not in ("", None):
            continue
        sym = r["symbol"]
        if r["exch_seg"] == "NSE":
            if not sym.endswith("-EQ"):
                continue
            sym = sym[:-3]
        out.append({"broker": "angel", "exchange": r["exch_seg"], "token": str(r["token"]),
                    "symbol": sym, "name": r["name"], "instrument_type": "EQ"})
    return pl.DataFrame(out, schema={"broker": pl.Utf8, "exchange": pl.Utf8, "token": pl.Utf8,
                                     "symbol": pl.Utf8, "name": pl.Utf8,
                                     "instrument_type": pl.Utf8})


def parse_screener_csv(content: bytes) -> pl.DataFrame:
    """Screener.in screen export. Mapped on the 'NSE Code' / 'BSE Code' columns the
    export includes; every other column is kept as a (field, value) pair.

    Never merged into point-in-time fundamentals: the export has no filing dates.
    """
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    header = [h.strip() for h in (reader.fieldnames or [])]
    if "NSE Code" not in header and "BSE Code" not in header:
        raise SchemaMismatch("Screener export needs an 'NSE Code' or 'BSE Code' column; add it "
                             "to the screen's columns before exporting")
    out = []
    for raw in reader:
        r = {k.strip(): (v or "").strip() for k, v in raw.items() if k}
        nse, bse = r.get("NSE Code") or None, r.get("BSE Code") or None
        for field, value in r.items():
            if field in ("NSE Code", "BSE Code", "S.No."):
                continue
            try:
                num = _num(value)
            except ValueError:
                num = None
            out.append({"nse_code": nse, "bse_code": bse, "field": field, "period_label": None,
                        "value_text": value, "value_num": num})
    return pl.DataFrame(out, schema={"nse_code": pl.Utf8, "bse_code": pl.Utf8, "field": pl.Utf8,
                                     "period_label": pl.Utf8, "value_text": pl.Utf8,
                                     "value_num": pl.Float64})


def parse_screener_excel(content: bytes, nse_code: str | None,
                         bse_code: str | None) -> pl.DataFrame:
    """Per-company Screener Excel export: the 'Data Sheet' tab as (field, period, value).

    Requires openpyxl (optional dependency group `screener`).
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - optional
        raise RuntimeError("install the 'screener' extra to import Excel exports") from exc
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "Data Sheet" not in wb.sheetnames:
        raise SchemaMismatch(f"no 'Data Sheet' tab; sheets: {wb.sheetnames}")
    out, periods = [], []
    for row in wb["Data Sheet"].iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        if label.lower() == "report date":
            periods = [str(c) if c is not None else None for c in row[1:]]
            continue
        if not periods:
            continue
        for period, value in zip(periods, row[1:], strict=False):
            if period is None or value is None:
                continue
            num = value if isinstance(value, (int, float)) else None
            out.append({"nse_code": nse_code, "bse_code": bse_code, "field": label,
                        "period_label": period, "value_text": str(value),
                        "value_num": float(num) if num is not None else None})
    return pl.DataFrame(out, schema={"nse_code": pl.Utf8, "bse_code": pl.Utf8, "field": pl.Utf8,
                                     "period_label": pl.Utf8, "value_text": pl.Utf8,
                                     "value_num": pl.Float64})
